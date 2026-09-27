from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

import os
import tempfile

from admin_panel.models import AdminUser
from courses.forms import StudentSignupForm

from courses.models import (
    Course,
    CourseCategory,
    CourseContent,
    Exam,
    ExamQuestion,
    Question,
    QuestionBank,
    Resource,
    StudentProfile,
)
from courses.resource_protection import docx_to_protected_html


class StudentSignupMobileTests(TestCase):
    def _valid_data(self, **overrides):
        data = {
            "full_name": "Riya Sharma",
            "username": "riya_s",
            "email": "riya@example.com",
            "country_code": "91",
            "mobile_number": "9876543210",
            "password": "LearningBanyanPass123",
        }
        data.update(overrides)
        return data

    def test_mobile_number_is_optional(self):
        form = StudentSignupForm(self._valid_data(mobile_number=""))
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        profile = StudentProfile.objects.get(user=user)
        self.assertIsNone(profile.mobile_number)

    def test_two_signups_can_omit_mobile(self):
        first = StudentSignupForm(self._valid_data(mobile_number=""))
        self.assertTrue(first.is_valid(), first.errors)
        first.save()
        second = StudentSignupForm(
            self._valid_data(
                username="other_user",
                email="other@example.com",
                mobile_number="",
            )
        )
        self.assertTrue(second.is_valid(), second.errors)
        second.save()
        self.assertEqual(
            StudentProfile.objects.filter(mobile_number__isnull=True).count(),
            2,
        )

    def test_invalid_mobile_is_rejected(self):
        form = StudentSignupForm(self._valid_data(mobile_number="12345"))
        self.assertFalse(form.is_valid())
        self.assertIn("mobile_number", form.errors)

    def test_signup_saves_mobile_number(self):
        form = StudentSignupForm(self._valid_data())
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        profile = StudentProfile.objects.get(user=user)
        self.assertEqual(profile.mobile_number, "+919876543210")

    def test_signup_view_allows_missing_mobile(self):
        response = self.client.post(
            reverse("courses:student_signup"), self._valid_data(mobile_number="")
        )
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="riya_s")
        self.assertIsNone(user.student_profile.mobile_number)

    def test_duplicate_mobile_is_rejected(self):
        existing = User.objects.create_user(
            username="existing",
            email="existing@example.com",
            password="LearningBanyanPass123",
        )
        StudentProfile.objects.create(user=existing, mobile_number="9876543210")
        form = StudentSignupForm(
            self._valid_data(username="other_user", email="other@example.com")
        )
        self.assertFalse(form.is_valid())
        self.assertIn("mobile_number", form.errors)

    def test_plus_91_mobile_matches_existing_ten_digit_number(self):
        existing = User.objects.create_user(
            username="existing",
            email="existing@example.com",
            password="LearningBanyanPass123",
        )
        StudentProfile.objects.create(user=existing, mobile_number="9876543210")
        form = StudentSignupForm(
            self._valid_data(
                username="other_user",
                email="other@example.com",
                mobile_number="+91 98765 43210",
            )
        )
        self.assertFalse(form.is_valid())
        self.assertIn("mobile_number", form.errors)

    def test_signup_view_creates_account_immediately(self):
        response = self.client.post(
            reverse("courses:student_signup"), self._valid_data()
        )
        self.assertRedirects(response, reverse("courses:student_exams"))
        user = User.objects.get(username="riya_s")
        profile = StudentProfile.objects.get(user=user)
        self.assertEqual(profile.mobile_number, "+919876543210")
        self.assertEqual(int(self.client.session.get("_auth_user_id")), user.pk)

    def test_signup_page_includes_country_code(self):
        response = self.client.get(reverse("courses:student_signup"))
        self.assertContains(response, "India (+91)")
        self.assertContains(response, 'name="country_code"')
        self.assertContains(response, "Create account")
        self.assertContains(response, "(optional)")
        self.assertNotContains(response, "Send OTP")
        self.assertNotContains(response, "one-time password")

    def test_other_country_code_is_stored(self):
        form = StudentSignupForm(
            self._valid_data(country_code="971", mobile_number="501234567")
        )
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        profile = StudentProfile.objects.get(user=user)
        self.assertEqual(profile.mobile_number, "+971501234567")


class AdminRegisteredUsersTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="admin",
            email="admin@example.com",
            password="admin123",
            first_name="Admin",
        )
        AdminUser.objects.create(user=self.admin)
        student = User.objects.create_user(
            username="student1",
            email="student1@example.com",
            password="LearningBanyanPass123",
            first_name="Aman",
            last_name="Singh",
        )
        StudentProfile.objects.create(user=student, mobile_number="9988776655")

    def test_dashboard_shows_registered_user_count(self):
        self.client.login(username="admin", password="admin123")
        response = self.client.get(reverse("admin_panel:dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Registered Users")
        self.assertContains(response, "1")
        self.assertEqual(response.context["total_registered_users"], 1)

    def test_manage_users_lists_students_and_mobile(self):
        self.client.login(username="admin", password="admin123")
        response = self.client.get(reverse("admin_panel:manage_users"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Aman Singh")
        self.assertContains(response, "student1@example.com")
        self.assertContains(response, "9988776655")
        self.assertContains(response, "1 registered")
        self.assertEqual(response.context["total_registered_users"], 1)

    def test_manage_users_search_filters_results(self):
        other = User.objects.create_user(
            username="other",
            email="other@example.com",
            password="LearningBanyanPass123",
            first_name="Neha",
        )
        StudentProfile.objects.create(user=other, mobile_number="9111222333")
        self.client.login(username="admin", password="admin123")
        response = self.client.get(reverse("admin_panel:manage_users"), {"q": "9988776655"})
        self.assertContains(response, "Aman Singh")
        self.assertNotContains(response, "Neha")
        self.assertEqual(response.context["total_registered_users"], 2)
        self.assertEqual(response.context["filtered_count"], 1)

    def test_students_cannot_open_registered_users(self):
        self.client.login(username="student1", password="LearningBanyanPass123")
        response = self.client.get(reverse("admin_panel:manage_users"), follow=True)
        self.assertNotEqual(response.request["PATH_INFO"], "/admin-panel/registered-users/")

    def test_dashboard_shows_currently_logged_in_admin(self):
        self.client.login(username="admin", password="admin123")
        response = self.client.get(reverse("admin_panel:dashboard"))
        self.assertContains(response, "Currently logged in")
        self.assertContains(response, "@admin")

    def test_manage_users_shows_currently_logged_in_admin(self):
        self.client.login(username="admin", password="admin123")
        response = self.client.get(reverse("admin_panel:manage_users"))
        self.assertContains(response, "Currently logged in")
        self.assertContains(response, "@admin")


class CurrentLoginDisplayTests(TestCase):
    def test_header_shows_logged_in_student(self):
        student = User.objects.create_user(
            username="student1",
            email="student1@example.com",
            password="LearningBanyanPass123",
            first_name="Aman",
            last_name="Singh",
        )
        StudentProfile.objects.create(user=student, mobile_number="9988776655")
        self.client.login(username="student1", password="LearningBanyanPass123")
        response = self.client.get(reverse("courses:home"))
        self.assertContains(response, "Logged in as")
        self.assertContains(response, "Aman Singh")
        self.assertContains(response, "student1")

    def test_header_hides_login_label_when_anonymous(self):
        response = self.client.get(reverse("courses:home"))
        self.assertNotContains(response, "Logged in as")


class ExamHintCheckboxTests(TestCase):
    def setUp(self):
        self.student = User.objects.create_user(
            username="hint_student",
            email="hint@example.com",
            password="LearningBanyanPass123",
        )
        StudentProfile.objects.create(user=self.student, mobile_number="+919000000001")
        category = CourseCategory.objects.create(name="Hint Maths")
        course = Course.objects.create(
            title="Hint Course",
            description="Practice",
            category=category,
        )
        bank = QuestionBank.objects.create(course=course, title="Bank")
        self.question = Question.objects.create(
            question_bank=bank,
            question_type="single_choice",
            question_text="What is 2+2?",
            option_a="4",
            correct_answer="A",
            hint="Add the numbers.",
        )
        self.exam = Exam.objects.create(
            name="Hint Paper",
            created_by=self.student,
            allow_hint=False,
        )
        ExamQuestion.objects.create(exam=self.exam, question=self.question, order=0)

    def test_hint_stays_hidden_until_selected(self):
        self.client.login(username="hint_student", password="LearningBanyanPass123")
        builder = self.client.get(
            reverse("courses:student_edit_exam", args=[self.exam.id])
        )
        self.assertContains(builder, "Need Hint")

        hidden = self.client.get(
            reverse("courses:student_practice_start", args=[self.exam.id])
        )
        self.assertNotContains(hidden, 'id="pqHintToggle"')
        self.assertNotContains(hidden, "Add the numbers.")

        self.exam.allow_hint = True
        self.exam.save(update_fields=["allow_hint"])
        shown = self.client.get(
            reverse("courses:student_practice_start", args=[self.exam.id])
        )
        self.assertContains(shown, 'id="pqHintToggle"')
        self.assertContains(shown, "Add the numbers.")


class DocxEquationPreviewTests(TestCase):
    def test_omml_equation_is_kept_in_preview(self):
        from docx import Document
        from docx.oxml import parse_xml

        math = (
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>x+1</m:t></m:r>"
            "</m:oMath>"
        )
        document = Document()
        paragraph = document.add_paragraph("Solve ")
        paragraph._p.append(parse_xml(math))

        handle = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        path = handle.name
        handle.close()
        try:
            document.save(path)

            class _File:
                def __init__(self, file_path):
                    self.path = file_path
                    self.name = os.path.basename(file_path)
                    self.size = os.path.getsize(file_path)

            html = docx_to_protected_html(_File(path))
        finally:
            os.unlink(path)

        self.assertIn("Solve", html)
        self.assertIn("x + 1", html)
        self.assertNotIn("rd-preview-error", html)

    def test_old_libreoffice_docx_still_opens(self):
        import zipfile

        from docx import Document

        document = Document()
        document.add_paragraph("Binomial expansion stays readable")
        handle = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        path = handle.name
        handle.close()
        try:
            document.save(path)
            with zipfile.ZipFile(path, "r") as zin:
                rels = zin.read("_rels/.rels").decode("utf-8")
                rels = rels.replace(
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
                    "http://purl.oclc.org/ooxml/officeDocument/relationships/officeDocument",
                )
                payload = {name: zin.read(name) for name in zin.namelist()}
            payload["_rels/.rels"] = rels.encode("utf-8")
            with zipfile.ZipFile(path, "w") as zout:
                for name, data in payload.items():
                    zout.writestr(name, data)

            class _File:
                def __init__(self, file_path):
                    self.path = file_path
                    self.name = os.path.basename(file_path)
                    self.size = os.path.getsize(file_path)

            html = docx_to_protected_html(_File(path))
        finally:
            os.unlink(path)

        self.assertIn("Binomial expansion stays readable", html)
        self.assertNotIn("rd-preview-error", html)


class CourseContentAdminTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="course_admin",
            email="course-admin@example.com",
            password="admin123",
        )
        AdminUser.objects.create(user=self.admin)
        self.category = CourseCategory.objects.create(name="CUET Maths")
        self.client.login(username="course_admin", password="admin123")

    def test_create_course_then_add_folder(self):
        response = self.client.post(
            reverse("admin_panel:create_course"),
            {
                "title": "CUET Mathematics 2026",
                "description": "Full course",
                "category": self.category.id,
                "level": "beginner",
                "price": "1499",
                "discount": "100",
                "subcategory": "CUET",
                "students_enrolled": "0",
                "rating": "0",
            },
        )
        course = Course.objects.get(title="CUET Mathematics 2026")
        self.assertRedirects(
            response, reverse("admin_panel:edit_course", args=[course.id])
        )
        self.assertEqual(course.sale_price, 1399)
        added = self.client.post(
            reverse("admin_panel:add_course_content", args=[course.id]),
            {"content_type": "folder", "content_title": "01 Relation and Function"},
        )
        self.assertRedirects(
            added, reverse("admin_panel:edit_course", args=[course.id])
        )
        item = CourseContent.objects.get(course=course)
        self.assertEqual(item.content_type, "folder")
        self.assertEqual(item.title, "01 Relation and Function")
        page = self.client.get(reverse("admin_panel:edit_course", args=[course.id]))
        self.assertContains(page, "Online Test/Quiz")
        self.assertContains(page, "Subjective Test")
        self.assertContains(page, "Practice Test")
        self.assertContains(page, "01 Relation and Function")
        self.assertContains(page, "Sub category")

    def test_create_page_content_options_open_fields(self):
        response = self.client.get(reverse("admin_panel:create_course"))
        self.assertContains(response, 'data-type="video"')
        self.assertContains(response, 'id="contentTitle"')
        self.assertContains(response, 'id="contentFile"')
        self.assertContains(response, "Course Content")
        self.assertContains(response, "Upload video")
        self.assertContains(response, "YouTube link")
        self.assertContains(response, 'id="contentVideoUrl"')
        self.assertNotContains(response, "Save the course first")

        created = self.client.post(
            reverse("admin_panel:create_course"),
            {
                "title": "CUET With Folder",
                "description": "Full course",
                "category": self.category.id,
                "level": "beginner",
                "students_enrolled": "0",
                "rating": "0",
                "content_type": "folder",
                "content_title": "Video Lectures",
            },
        )
        course = Course.objects.get(title="CUET With Folder")
        self.assertRedirects(
            created, reverse("admin_panel:edit_course", args=[course.id])
        )
        self.assertTrue(
            CourseContent.objects.filter(
                course=course, content_type="folder", title="Video Lectures"
            ).exists()
        )

    def test_video_requires_a_file(self):
        course = Course.objects.create(
            title="Video Course",
            description="Desc",
            category=self.category,
        )
        response = self.client.post(
            reverse("admin_panel:add_course_content", args=[course.id]),
            {"content_type": "video", "content_title": "Lecture 1"},
        )
        self.assertEqual(CourseContent.objects.filter(course=course).count(), 0)
        self.assertRedirects(
            response, reverse("admin_panel:edit_course", args=[course.id])
        )

    def test_video_accepts_a_youtube_link(self):
        course = Course.objects.create(
            title="YouTube Course",
            description="Desc",
            category=self.category,
        )
        response = self.client.post(
            reverse("admin_panel:add_course_content", args=[course.id]),
            {
                "content_type": "video",
                "content_title": "Lecture 1",
                "content_video_source": "youtube",
                "content_video_url": "https://youtu.be/dQw4w9WgXcQ",
            },
        )
        self.assertRedirects(
            response, reverse("admin_panel:edit_course", args=[course.id])
        )
        item = CourseContent.objects.get(course=course)
        self.assertEqual(item.video_url, "https://youtu.be/dQw4w9WgXcQ")
        self.assertFalse(item.file)
        self.assertEqual(
            item.youtube_embed_url(), "https://www.youtube.com/embed/dQw4w9WgXcQ"
        )
        page = self.client.get(reverse("courses:course_detail", args=[course.id]))
        self.assertContains(page, "https://www.youtube.com/embed/dQw4w9WgXcQ")

    def test_video_rejects_a_non_youtube_link(self):
        course = Course.objects.create(
            title="Bad Link Course",
            description="Desc",
            category=self.category,
        )
        self.client.post(
            reverse("admin_panel:add_course_content", args=[course.id]),
            {
                "content_type": "video",
                "content_title": "Lecture 1",
                "content_video_source": "youtube",
                "content_video_url": "https://example.com/watch",
            },
        )
        self.assertEqual(CourseContent.objects.filter(course=course).count(), 0)


class WordResourcePreviewTests(TestCase):
    def _resource(self, name, content=b"not-a-real-doc"):
        from django.core.files.uploadedfile import SimpleUploadedFile

        return Resource.objects.create(
            title="Algebra notes",
            resource_type="other",
            file=SimpleUploadedFile(name, content),
        )

    def test_doc_opens_in_pdf_reader_when_layout_export_works(self):
        from unittest.mock import patch

        resource = self._resource("notes.doc")
        with patch(
            "courses.resource_protection.ensure_word_preview_pdf",
            return_value="resource_previews/notes.pdf",
        ):
            page = self.client.get(
                reverse("courses:resource_detail", args=[resource.pk])
            )
        self.assertContains(page, "rdPdfWrap")
        self.assertContains(page, "Opening the Word file")
        self.assertContains(page, "opacity: 0.035")
        self.assertNotContains(page, "Legacy .doc")

    def test_doc_shows_document_text_when_pdf_export_fails(self):
        from unittest.mock import patch

        resource = self._resource("notes.doc")
        with patch(
            "courses.resource_protection.ensure_word_preview_pdf",
            return_value="",
        ), patch(
            "courses.resource_protection.legacy_doc_to_protected_html",
            return_value="<p class='rd-para'>If vertices of a triangle</p>",
        ):
            page = self.client.get(
                reverse("courses:resource_detail", args=[resource.pk])
            )
        self.assertContains(page, "If vertices of a triangle")
        self.assertNotContains(page, "rdPdfWrap")
        self.assertNotContains(page, "Legacy .doc")

    def test_stream_serves_word_file_as_pdf(self):
        import io
        from unittest.mock import patch

        resource = self._resource("notes.docx", b"docx-bytes")
        with patch(
            "courses.resource_protection.open_word_preview_pdf",
            return_value=io.BytesIO(b"%PDF-1.4 preview"),
        ):
            response = self.client.get(
                reverse("courses:resource_file_stream", args=[resource.pk]),
                HTTP_SEC_FETCH_DEST="empty",
                HTTP_SEC_FETCH_MODE="cors",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        body = b"".join(response.streaming_content)
        self.assertTrue(body.startswith(b"%PDF"))

    def test_word_preview_pdf_is_cached(self):
        import io
        import os
        from unittest.mock import patch

        from django.core.files.storage import default_storage

        class _Upload:
            name = "resources/cached-notes.doc"
            size = 4

            def __init__(self):
                self.path = ""

            def open(self, mode="rb"):
                self._buf = io.BytesIO(b"doc!")
                return self

            def read(self, *args):
                return self._buf.read(*args)

            def close(self):
                return None

        def _write_pdf(src_path, outdir):
            out = os.path.join(outdir, "out.pdf")
            with open(out, "wb") as handle:
                handle.write(b"%PDF-1.4 cached")
            return out

        upload = _Upload()
        with patch(
            "courses.resource_protection._convert_word_source_to_pdf",
            side_effect=_write_pdf,
        ) as convert:
            from courses.resource_protection import ensure_word_preview_pdf

            first = ensure_word_preview_pdf(upload)
            second = ensure_word_preview_pdf(upload)
        self.assertTrue(first.endswith(".pdf"))
        self.assertEqual(first, second)
        self.assertEqual(convert.call_count, 1)
        self.assertTrue(default_storage.exists(first))
        default_storage.delete(first)

    def test_preview_pdf_url_is_blocked(self):
        from django.http import HttpResponse
        from django.test import RequestFactory

        from courses.middleware import InlineMediaMiddleware

        middleware = InlineMediaMiddleware(lambda request: HttpResponse("file"))
        request = RequestFactory().get("/media/resource_previews/abc.pdf")
        response = middleware(request)
        self.assertEqual(response.status_code, 403)
